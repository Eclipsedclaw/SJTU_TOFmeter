#!/usr/bin/env python3
"""Convert an EXPACS 3.x workbook into the spectrum table read by CosmicRaySource.

    python3 expacs_to_table.py EXPACS303-eng.xlsx [output.txt]

Default output: ../data/expacs_spectrum.txt next to this script. The script reads the
values cached in the "Main" sheet, i.e. the spectra EXPACS computed for the conditions
entered there: set the conditions in Excel, let it recalculate and save, then run this.
Needs openpyxl. EXPACS: T. Sato, PLoS ONE 10(12): e0144679 (2015),
https://phits.jaea.go.jp/expacs/
"""
import os
import sys

import openpyxl

# EXPACS column label -> Geant4 particle name ("<symbol><A>" for ions)
SPECIES = {
    "Neutron": "neutron", "Proton": "proton", "He ion": "alpha",
    "Positive muon": "mu+", "Negative muon": "mu-", "Electron": "e-",
    "Positron": "e+", "Photon": "gamma",
    "Li ion": "Li7", "Be ion": "Be9", "B ion": "B11", "C ion": "C12", "N ion": "N14",
    "O ion": "O16", "F ion": "F19", "Ne ion": "Ne20", "Na ion": "Na23", "Mg ion": "Mg24",
    "Al ion": "Al27", "Si ion": "Si28", "P ion": "P31", "S ion": "S32", "Cl ion": "Cl35",
    "Ar ion": "Ar40", "K ion": "K39", "Ca ion": "Ca40", "Sc ion": "Sc45", "Ti ion": "Ti48",
    "V ion": "V51", "Cr ion": "Cr52", "Mn ion": "Mn55", "Fe ion": "Fe56", "Co ion": "Co59",
    "Ni ion": "Ni58",
}
HEADER_ROW, FIRST_ROW, ENERGY_COL, FIRST_SPECIES_COL = 34, 35, 4, 5  # Main sheet layout


def number(value):
    return float(value) if isinstance(value, (int, float)) and value > 0 else 0.0


def conditions(ws):
    """Input conditions of the Main sheet, decoded with the option codes of EXPACS 3.03."""
    altitude_unit = {1: "g/cm2 (depth)", 2: "km", 3: "ft", 4: "hPa"}.get(ws["C7"].value, "?")
    if ws["C8"].value == 2:
        location = f"latitude {ws['B8'].value} deg, longitude {ws['B9'].value} deg"
    else:
        location = f"cut-off rigidity {ws['B8'].value} GV given"
    solar = {1: f"W value {ws['B10'].value} given",
             2: f"date {ws['B10'].value}-{ws['B11'].value:02d}-{ws['B12'].value:02d}",
             3: f"Oulu neutron monitor {ws['B10'].value} counts/min"}.get(ws["C10"].value, "?")
    environment = {1: "ground", 2: "aircraft (pilot)", 3: "aircraft (cabin)", 4: "none"}.get(
        ws["B13"].value, "?")
    return [
        f"altitude {ws['B7'].value} {altitude_unit} -> atmospheric depth {ws['B24'].value} g/cm2",
        f"{location} -> cut-off rigidity {ws['B25'].value} GV",
        f"{solar} -> solar activity W = {ws['B26'].value}",
        f"surroundings: {environment}, local effect parameter {ws['B27'].value} ({ws['C27'].value})",
    ]


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    workbook = sys.argv[1]
    default = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "expacs_spectrum.txt")
    output = sys.argv[2] if len(sys.argv) > 2 else default

    ws = openpyxl.load_workbook(workbook, data_only=True)["Main"]
    if ws["C16"].value != 1:
        sys.exit("Set 'Output flux unit' (Main!C16) to 1, phi in /cm2/s/(MeV/n), and recalculate")

    columns, names = [], []
    col = FIRST_SPECIES_COL
    while ws.cell(HEADER_ROW, col).value:
        label = ws.cell(HEADER_ROW, col).value
        if label not in SPECIES:
            sys.exit(f"Unknown EXPACS column '{label}'")
        columns.append(col)
        names.append(SPECIES[label])
        col += 1
    rows = []
    row = FIRST_ROW
    while isinstance(ws.cell(row, ENERGY_COL).value, (int, float)):
        rows.append([float(ws.cell(row, ENERGY_COL).value)] + [number(ws.cell(row, c).value) for c in columns])
        row += 1

    os.makedirs(os.path.dirname(os.path.abspath(output)), exist_ok=True)
    with open(output, "w") as out:
        out.write(f"# {ws['B1'].value} (PARMA model, T. Sato, PLoS ONE 10(12): e0144679, 2015)\n")
        out.write(f"# converted from {os.path.basename(workbook)} by scripts/expacs_to_table.py\n")
        for line in conditions(ws):
            out.write(f"# {line}\n")
        out.write("# phi: omnidirectional fluence rate in 1/(cm2 s (MeV/n)); E: kinetic energy per nucleon (MeV)\n")
        out.write("# columns: E " + " ".join(names) + "\n")
        for values in rows:
            out.write(" ".join(f"{v:.6e}" for v in values) + "\n")
    print(f"wrote {len(rows)} energies x {len(names)} species to {output}")


if __name__ == "__main__":
    main()
